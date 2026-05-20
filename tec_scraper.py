#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TecAlliance OE Cross-Reference Scraper
=======================================
אוסף OE numbers ורכבים תואמים מ-TecAlliance לכל מוצר בקובץ Excel.

הרצה:
  pip install playwright openpyxl
  python -m playwright install chromium
  python tec_scraper.py
"""

import sys
import time
import json
import re
import os
from pathlib import Path
from collections import defaultdict

# ─── הגדרות - שנה לפי הצורך ────────────────────────────────────────────────
EXCEL_INPUT  = r"Copy_of_____________________.xlsx"   # נתיב לקובץ המקורי
EXCEL_OUTPUT = r"output_with_oe_numbers.xlsx"          # נתיב לקובץ הפלט
CHECKPOINT   = "tec_checkpoint.json"                   # קובץ שמירת התקדמות
BASE_URL     = "https://web.tecalliance.net/mvspareparts/qa/"
DELAY_SEC    = 1.5    # שניות בין חיפושים
HEADLESS     = False  # False = הדפדפן נראה; True = ברקע
DEBUG_SHOTS  = True   # שמור screenshots לאיתור באגים
# ────────────────────────────────────────────────────────────────────────────


def _install():
    import subprocess
    for pkg in ["playwright", "openpyxl"]:
        try:
            __import__(pkg.replace("-", "_"))
        except ImportError:
            print(f"Installing {pkg}...")
            subprocess.check_call([sys.executable, "-m", "pip", "install", pkg, "-q"])
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            exe = p.chromium.executable_path
            if not Path(exe).exists():
                raise FileNotFoundError
    except Exception:
        print("Installing Chromium browser...")
        import subprocess
        subprocess.check_call([sys.executable, "-m", "playwright", "install", "chromium"])


_install()

from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


# ═══════════════════════════════════════════════════════════
#  1. טעינת Excel
# ═══════════════════════════════════════════════════════════

def load_products(path: str) -> dict:
    """קורא גיליון1 ומחזיר dict לפי SAP ID."""
    wb = openpyxl.load_workbook(path)
    ws = wb["גיליון1"]
    products = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        sap_id     = str(int(row[0])) if isinstance(row[0], float) else str(row[0])
        sku_code   = str(row[1]).strip() if row[1] else ""
        oem_code   = str(row[2]).strip() if row[2] else ""
        description= str(row[3] or "")
        sub_cat    = str(row[4] or "")
        brand      = str(row[5] or "")

        if sap_id not in products:
            products[sap_id] = {
                "description": description,
                "sub_cat": sub_cat,
                "brand": brand,
                "sku_codes": [],
                "oem_codes": [],
            }
        p = products[sap_id]
        if sku_code and sku_code not in p["sku_codes"]:
            p["sku_codes"].append(sku_code)
        if oem_code and oem_code not in p["oem_codes"]:
            p["oem_codes"].append(oem_code)

    print(f"[Excel] Loaded {len(products):,} unique SAP IDs")
    return products


def unique_oem_map(products: dict) -> dict:
    """מיפוי: oem_code → [sap_ids]"""
    m = defaultdict(list)
    for sap_id, p in products.items():
        for oem in p["oem_codes"]:
            if oem and oem not in m[oem]:
                m[oem].append(sap_id)
    return dict(m)


# ═══════════════════════════════════════════════════════════
#  2. Playwright scraper
# ═══════════════════════════════════════════════════════════

class TecScraper:
    SEARCH_SELECTORS = [
        'input[placeholder*="Search according"]',
        'input[placeholder*="Search"]',
        'input[placeholder*="number"]',
        'input[type="search"]',
        '[class*="SearchInput"] input',
        '[class*="search-input"] input',
        '[class*="searchInput"] input',
        'input.search',
    ]
    RESULT_SELECTORS = [
        '[class*="ArticleList"] [class*="ArticleItem"]:first-child',
        '[class*="article-list"] [class*="article-item"]:first-child',
        '[class*="ResultList"] li:first-child a',
        '[class*="result-list"] li:first-child a',
        '[class*="articles"] > div:first-child',
        'article:first-child',
    ]
    OE_NUMBER_SELECTORS = [
        '[class*="OeNumber"]',
        '[class*="oe-number"]',
        '[class*="oeNumber"]',
        '[class*="ArticleNumber"]:not([class*="Header"])',
        'td[class*="oe"]',
        '[data-label="OE number"]',
        'td:nth-child(1)[class*="number"]',
    ]
    REF_INFO_SELECTORS = [
        '[class*="ReferenceInfo"]',
        '[class*="reference-info"]',
        '[class*="referenceInfo"]',
        '[data-label="Reference info"]',
        'td:nth-child(2)[class*="ref"]',
    ]

    def __init__(self, page, debug_shots: bool = False):
        self.page = page
        self.debug_shots = debug_shots
        self._shot_idx = 0
        self._captured_oe: list = []   # data intercepted from API
        page.on("response", self._on_response)

    # ── API interception ────────────────────────────────────
    def _on_response(self, response):
        try:
            if response.status != 200:
                return
            ct = response.headers.get("content-type", "")
            if "json" not in ct:
                return
            url = response.url.lower()
            if not any(k in url for k in ["article", "part", "oe", "reference", "cross", "search"]):
                return
            body = response.json()
            rows = self._extract_oe_from_json(body)
            if rows:
                self._captured_oe.extend(rows)
        except Exception:
            pass

    def _extract_oe_from_json(self, obj, depth=0) -> list:
        """Recursively hunt for OE-number + reference pairs in JSON."""
        if depth > 12:
            return []
        results = []

        if isinstance(obj, dict):
            # Direct pattern: {oeNumber: "...", referenceInfo: "..."}
            oe_keys  = [k for k in obj if re.search(r"oe.?num|oem.?num|article.?num", k, re.I)]
            ref_keys = [k for k in obj if re.search(r"ref|vehicle|brand|manu|descr", k, re.I)]
            if oe_keys:
                oe_val  = str(obj[oe_keys[0]]).strip()
                ref_val = str(obj[ref_keys[0]]).strip() if ref_keys else ""
                if oe_val and len(oe_val) >= 3:
                    results.append({"oe": oe_val, "ref": ref_val})
            for v in obj.values():
                results.extend(self._extract_oe_from_json(v, depth + 1))

        elif isinstance(obj, list):
            for item in obj:
                results.extend(self._extract_oe_from_json(item, depth + 1))

        return results

    # ── screenshot helper ───────────────────────────────────
    def _shot(self, label: str):
        if not self.debug_shots:
            return
        fname = f"debug_{self._shot_idx:03d}_{label}.png"
        try:
            self.page.screenshot(path=fname)
        except Exception:
            pass
        self._shot_idx += 1

    # ── find element with fallback list ────────────────────
    def _find(self, selectors, timeout=4000):
        for sel in selectors:
            try:
                el = self.page.wait_for_selector(sel, timeout=timeout)
                if el:
                    return el
            except Exception:
                pass
        return None

    def _find_all(self, selectors):
        for sel in selectors:
            els = self.page.query_selector_all(sel)
            if els:
                return els
        return []

    # ── main search ─────────────────────────────────────────
    def search(self, oem_code: str) -> list:
        """
        Returns list of {"oe": str, "ref": str} dicts.
        """
        self._captured_oe.clear()

        try:
            # Step 1: type in search box
            search_el = self._find(self.SEARCH_SELECTORS)
            if not search_el:
                print(f"  ⚠️  Search box not found for: {oem_code}")
                return []

            search_el.click(click_count=3)
            search_el.fill("")
            search_el.type(oem_code, delay=40)
            time.sleep(0.6)

            # Step 2: autocomplete or Enter
            suggestion = None
            for sel in [
                '[class*="Suggestion"]:first-child',
                '[class*="suggestion"]:first-child',
                '[class*="autocomplete"] li:first-child',
                '[class*="dropdown"] [class*="item"]:first-child',
            ]:
                try:
                    suggestion = self.page.wait_for_selector(sel, timeout=1500)
                    break
                except Exception:
                    pass

            if suggestion:
                suggestion.click()
            else:
                search_el.press("Enter")

            self.page.wait_for_load_state("networkidle", timeout=15000)
            time.sleep(0.8)
            self._shot("after_search")

            # Step 3: click first result
            result_el = self._find(self.RESULT_SELECTORS, timeout=5000)
            if not result_el:
                # fallback: try clicking any visible article link
                result_el = self.page.query_selector("a[href*='/parts/']")
            if not result_el:
                print(f"  ⚠️  No result for: {oem_code}")
                return []

            result_el.click()
            self.page.wait_for_load_state("networkidle", timeout=15000)
            time.sleep(1.0)
            self._shot("detail_page")

            # Step 4: try API-captured data first
            if self._captured_oe:
                return self._captured_oe[:]

            # Step 5: DOM fallback
            return self._scrape_oe_table()

        except PWTimeout:
            print(f"  ⏱️  Timeout: {oem_code}")
            return []
        except Exception as e:
            print(f"  ❌  Error ({oem_code}): {e}")
            return []

    def _scrape_oe_table(self) -> list:
        """Scrape OE number table directly from DOM."""
        results = []

        # Try to find a table with OE number rows
        # Pattern 1: rows with paired OE + reference cells
        rows = self.page.query_selector_all("tr")
        for row in rows:
            cells = row.query_selector_all("td")
            if len(cells) >= 2:
                oe  = cells[0].inner_text().strip()
                ref = cells[1].inner_text().strip()
                if oe and len(oe) >= 3 and oe.lower() not in ("oe number", "number", ""):
                    results.append({"oe": oe, "ref": ref})

        if results:
            return results

        # Pattern 2: separate OE and ref element lists
        oe_els  = self._find_all(self.OE_NUMBER_SELECTORS)
        ref_els = self._find_all(self.REF_INFO_SELECTORS)
        for i, el in enumerate(oe_els):
            oe  = el.inner_text().strip()
            ref = ref_els[i].inner_text().strip() if i < len(ref_els) else ""
            if oe and len(oe) >= 3:
                results.append({"oe": oe, "ref": ref})

        return results


# ═══════════════════════════════════════════════════════════
#  3. Checkpoint helpers
# ═══════════════════════════════════════════════════════════

def load_checkpoint(path):
    if Path(path).exists():
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        print(f"[Checkpoint] Resumed: {len(data):,} OEM codes done")
        return data
    return {}


def save_checkpoint(data, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# ═══════════════════════════════════════════════════════════
#  4. יצירת Excel
# ═══════════════════════════════════════════════════════════

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _header_row(ws, cols):
    ws.append(cols)
    for cell in ws[1]:
        cell.font  = HEADER_FONT
        cell.fill  = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def create_excel(products: dict, oem_results: dict, output_path: str):
    """
    Sheet 1 - מוצרים:
      SAP ID | תיאור | קטגוריה | מותג | SKU_1..N | OEM_1..N

    Sheet 2 - צליבת מק"טים (VLOOKUP by OEM code):
      SAP ID | תיאור | מק"ט חיפוש | OE Number | Reference Info
    """
    wb = openpyxl.Workbook()

    # ── Sheet 1: Products ───────────────────────────────────
    ws1 = wb.active
    ws1.title = "מוצרים"

    max_sku = max((len(p["sku_codes"]) for p in products.values()), default=1)
    max_oem = max((len(p["oem_codes"]) for p in products.values()), default=1)

    h1 = ["SAP ID", "תיאור", "קטגוריה", "מותג"]
    h1 += [f'SKU {i+1}' for i in range(max_sku)]
    h1 += [f'מק"ט OEM {i+1}' for i in range(max_oem)]
    _header_row(ws1, h1)

    for sap_id, p in products.items():
        row = [sap_id, p["description"], p["sub_cat"], p["brand"]]
        row += p["sku_codes"] + [""] * (max_sku - len(p["sku_codes"]))
        row += p["oem_codes"] + [""] * (max_oem - len(p["oem_codes"]))
        ws1.append(row)

    ws1.freeze_panes = "A2"

    # ── Sheet 2: Cross-reference ────────────────────────────
    ws2 = wb.create_sheet('צליבת מק"טים')
    _header_row(ws2, ["SAP ID", "תיאור", 'מק"ט חיפוש (OEM)', "OE Number", "Reference Info"])

    for sap_id, p in products.items():
        for oem_code in p["oem_codes"]:
            rows = oem_results.get(oem_code, [])
            if rows:
                for entry in rows:
                    ws2.append([
                        sap_id,
                        p["description"],
                        oem_code,
                        entry.get("oe", ""),
                        entry.get("ref", ""),
                    ])
            else:
                ws2.append([sap_id, p["description"], oem_code, "", ""])

    ws2.freeze_panes = "A2"

    # Auto-width
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=0)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    wb.save(output_path)
    print(f"\n✅  Excel saved → {output_path}")


# ═══════════════════════════════════════════════════════════
#  5. Main
# ═══════════════════════════════════════════════════════════

def main():
    print("=" * 58)
    print("  TecAlliance OE Cross-Reference Collector")
    print("=" * 58)

    # Resolve paths
    excel_input = EXCEL_INPUT
    if not Path(excel_input).exists():
        excel_input = input("\nנתיב לקובץ Excel: ").strip().strip('"').strip("'")
        if not Path(excel_input).exists():
            print(f"❌  File not found: {excel_input}")
            sys.exit(1)

    excel_output = EXCEL_OUTPUT
    if excel_output == r"output_with_oe_numbers.xlsx":
        p = Path(excel_input)
        excel_output = str(p.parent / (p.stem + "_OE_cross_ref.xlsx"))

    print(f"\nInput : {excel_input}")
    print(f"Output: {excel_output}")

    # Load
    products   = load_products(excel_input)
    oem_map    = unique_oem_map(products)
    checkpoint = load_checkpoint(CHECKPOINT)
    oem_results = dict(checkpoint)  # oem_code → [{"oe":..., "ref":...}]

    remaining = {k: v for k, v in oem_map.items() if k not in oem_results}
    print(f"\nUnique OEM codes : {len(oem_map):,}")
    print(f"Already done     : {len(oem_results):,}")
    print(f"Remaining        : {len(remaining):,}")

    if not remaining:
        print("\nAll done! Creating Excel...")
        create_excel(products, oem_results, excel_output)
        return

    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=HEADLESS, slow_mo=30)
        page    = browser.new_page()
        page.goto(BASE_URL)

        print("\n" + "=" * 58)
        print("  התחבר לאתר TecAlliance בדפדפן שנפתח.")
        print("  לאחר שנכנסת לחלוטין, חזור לחלון הזה ולחץ Enter.")
        print("=" * 58)
        input("\n  > ")

        scraper = TecScraper(page, debug_shots=DEBUG_SHOTS)
        total   = len(remaining)
        done    = 0

        for oem_code, sap_ids in remaining.items():
            done += 1
            preview_saps = ", ".join(sap_ids[:3])
            print(f"[{done}/{total}]  {oem_code:<20}  SAP: {preview_saps}")

            rows = scraper.search(oem_code)
            oem_results[oem_code] = rows

            n_oe  = len(rows)
            n_ref = sum(1 for r in rows if r.get("ref"))
            print(f"         → {n_oe} OE numbers, {n_ref} with reference info")

            if done % 100 == 0:
                save_checkpoint(oem_results, CHECKPOINT)
                print(f"  💾  Checkpoint saved [{done}/{total}]")

            time.sleep(DELAY_SEC)

        browser.close()

    save_checkpoint(oem_results, CHECKPOINT)
    print("\nCreating Excel output...")
    create_excel(products, oem_results, excel_output)
    print("\n✅  All done!")


if __name__ == "__main__":
    main()
